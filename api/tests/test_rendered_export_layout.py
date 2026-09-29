import base64
import unittest
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw
from pypdf import PdfReader
from pptx import Presentation

from app.exports import report_pdf, report_pptx


def sample_rendered_page(name="Executive Overview"):
    image = Image.new("RGB", (1600, 900), "#f8fafc")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((70, 70, 1530, 830), radius=24, fill="#ffffff", outline="#cbd5e1", width=3)
    draw.text((120, 115), "Revenue Dashboard", fill="#0f172a")
    for index, height in enumerate((240, 390, 520, 680, 440)):
        left = 150 + index * 260
        draw.rectangle((left, 760 - height, left + 150, 760), fill="#2563eb")
    stream = BytesIO()
    image.save(stream, format="PNG")
    encoded = base64.b64encode(stream.getvalue()).decode("ascii")
    return {
        "name": name,
        "image": f"data:image/png;base64,{encoded}",
        "insights": [],
        "narratives": [
            {"title": "Sales by product", "points": ["Highest sales: Laptop (86.0K)", "Lowest sales: Accessories (3.7K)"]},
            {"title": "Monthly trend", "points": ["Revenue increased 18.4% from January to June"]},
        ],
    }


class RenderedExportLayoutTests(unittest.TestCase):
    def setUp(self):
        self.project = {"report": {"name": "Quarterly Sales Report"}}
        self.pages = [sample_rendered_page(), sample_rendered_page("Regional Detail")]

    def test_pdf_has_header_date_divider_content_and_page_footer(self):
        document = PdfReader(BytesIO(report_pdf(self.project, self.pages)))
        self.assertEqual(len(document.pages), 2)
        for page_number, page in enumerate(document.pages, 1):
            text = page.extract_text()
            self.assertIn("Quarterly Sales Report", text)
            self.assertIn("Generated", text)
            self.assertIn(f"Page {page_number} of 2", text)
            self.assertTrue(page.images)

    def test_powerpoint_has_one_content_slide_per_report_page_with_header(self):
        deck = Presentation(BytesIO(report_pptx(self.project, self.pages)))
        self.assertEqual(len(deck.slides), 2)
        for page_number, slide in enumerate(deck.slides, 1):
            text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame"))
            self.assertIn("Quarterly Sales Report", text)
            self.assertIn("Generated", text)
            self.assertIn(f"Page {page_number} of 2", text)
            self.assertIn("SMART NARRATIVE", text)
            self.assertIn("Highest sales: Laptop (86.0K)", text)
            self.assertIn("Lowest sales: Accessories (3.7K)", text)
            self.assertTrue(any(shape.shape_type == 13 for shape in slide.shapes), "Captured report image is missing")

    def test_browser_export_waits_for_visual_data_instead_of_a_fixed_delay(self):
        source = (Path(__file__).parents[2] / "src" / "v11" / "PublishedViewer.tsx").read_text(encoding="utf-8")
        self.assertIn("waitForExportPage", source)
        self.assertIn("data-export-state", source)
        self.assertIn("data-page-id", source)
        self.assertIn("data-export-narrative", source)
        self.assertIn("Highest", source)
        self.assertIn("Lowest", source)
        self.assertNotIn("waitForRenderedPage", source)


if __name__ == "__main__":
    unittest.main()
