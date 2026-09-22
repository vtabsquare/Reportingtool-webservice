from __future__ import annotations
import base64, io, json, os, re, tempfile, urllib.parse, urllib.request, ipaddress, socket
import xml.etree.ElementTree as ET
from pathlib import Path
import pandas as pd
from .local_engine import metadata as local_metadata, write_dataframe, import_path

def safe_table_name(name:str)->str:
    stem=Path(name).stem
    stem=re.sub(r'[^A-Za-z0-9_]+','_',stem).strip('_') or 'ImportedData'
    if stem[0].isdigit():stem='T_'+stem
    return 'Imported_'+stem[:70]

def friendly_table_name(name:str)->str:
    """Return a concise semantic/display name without changing the physical table id.

    File imports keep a stable physical name such as ``Imported_SalesData_1000``.
    Authoring panes should show the business-facing name instead.  Large trailing
    record-count suffixes are removed only when they look like generated test/file
    counts (three or more digits), while normal business digits such as FY24 remain.
    """
    raw=Path(str(name or 'Table')).stem
    raw=re.sub(r'^(Imported_|ETL_)+','',raw,flags=re.I)
    raw=re.sub(r'[^A-Za-z0-9_ ]+','_',raw)
    raw=re.sub(r'[_ ]+\d{3,}$','',raw).strip(' _') or 'Table'
    # Keep identifiers compact and predictable for DAX/model display while preserving case.
    raw=re.sub(r'_+','_',raw)
    if raw[0].isdigit():raw='T_'+raw
    return raw[:48]

def demo_metadata():
    return local_metadata()

def workbook_sheets(path:Path):
    """Return workbook sheet info (name and column count) without importing data."""
    ext=path.suffix.lower()
    if ext not in ('.xlsx','.xls'):
        raise ValueError('Sheet discovery is available only for XLSX and XLS files.')
    engine='xlrd' if ext=='.xls' else 'openpyxl'
    out = []
    with pd.ExcelFile(path,engine=engine) as book:
        for x in book.sheet_names:
            try:
                df = pd.read_excel(book, sheet_name=x, nrows=0)
                out.append({'name': str(x), 'cols': len(df.columns), 'fields': [str(c) for c in df.columns]})
            except Exception:
                out.append({'name': str(x), 'cols': 0, 'fields': []})
    return out

def _json_frame(path:Path):
    """Normalize JSON/JSONL records, including nested objects and arrays."""
    if path.suffix.lower()=='.jsonl':
        records=[]
        with path.open('r',encoding='utf-8-sig') as handle:
            for line in handle:
                if line.strip(): records.append(json.loads(line))
    else:
        with path.open('r',encoding='utf-8-sig') as handle:
            payload=json.load(handle)
        if isinstance(payload,list): records=payload
        elif isinstance(payload,dict):
            list_values=[v for v in payload.values() if isinstance(v,list)]
            records=list_values[0] if len(list_values)==1 else [payload]
        else: records=[{'value':payload}]
    frame=pd.json_normalize(records,sep='.')
    for column in frame.columns:
        frame[column]=frame[column].map(lambda value:json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else value)
    return frame

def _xml_frame(path:Path):
    """Flatten common repeated-record XML while retaining attributes."""
    root=ET.parse(path).getroot()
    children=list(root)
    if not children:
        return pd.DataFrame([{'value':(root.text or '').strip(),**{f'@{k}':v for k,v in root.attrib.items()}}])
    counts={}
    for child in children:
        tag=child.tag.split('}')[-1];counts[tag]=counts.get(tag,0)+1
    repeated=max(counts,key=counts.get)
    records=[child for child in children if child.tag.split('}')[-1]==repeated]
    if len(records)==1 and list(records[0]): records=list(records[0])
    def flatten(node,prefix='',out=None):
        out={} if out is None else out
        for key,value in node.attrib.items():out[f'{prefix}@{key}']=value
        grouped={}
        for child in list(node):grouped.setdefault(child.tag.split('}')[-1],[]).append(child)
        for tag,nodes in grouped.items():
            key=f'{prefix}{tag}'
            if len(nodes)>1:
                out[key]=json.dumps([flatten(n,'',{}) if list(n) or n.attrib else (n.text or '').strip() for n in nodes],ensure_ascii=False)
            else:
                child=nodes[0]
                if list(child) or child.attrib:flatten(child,key+'.',out)
                else:out[key]=(child.text or '').strip()
        if not grouped and (node.text or '').strip():out[prefix.rstrip('.') or 'value']=(node.text or '').strip()
        return out
    return pd.DataFrame([flatten(node) for node in records])

def import_file_path(filename:str,path:Path,sheet:str|None=None,source_bytes:int|None=None):
    ext=Path(filename).suffix.lower()
    table=safe_table_name(f"{Path(filename).stem}_{sheet}" if sheet else filename)
    if ext in ('.csv','.tsv','.txt','.parquet'):
        rows=import_path(path,table,source_bytes)
        meta=next(x for x in local_metadata() if x['name']==table)
        return {'ok':True,'table':table,'rows':rows,'columns':[c['name'] for c in meta['columns']],'metadata':meta,'storage':meta.get('storage'),'sourceType':'file','displayName':friendly_table_name(table)}
    if ext in ('.json','.jsonl'):
        df=_json_frame(path)
        rows=write_dataframe(df,table,source_bytes or path.stat().st_size,'managed')
        meta=next(x for x in local_metadata() if x['name']==table)
        return {'ok':True,'table':table,'rows':rows,'columns':[str(c) for c in df.columns],'metadata':meta,'storage':meta.get('storage'),'sourceType':'file','displayName':friendly_table_name(table)}
    if ext in ('.xlsx','.xls'):
        engine='xlrd' if ext=='.xls' else 'openpyxl'
        sheets_info=workbook_sheets(path)
        sheet_names=[s['name'] for s in sheets_info]
        selected=sheet if sheet in sheet_names else sheet_names[0]
        # Each worksheet must have its own durable physical id.  Older packaged
        # backends reused the workbook id and silently overwrote earlier sheets.
        table=safe_table_name(f"{Path(filename).stem}_{selected}")
        df=pd.read_excel(path,sheet_name=selected,engine=engine)
        rows=write_dataframe(df,table,source_bytes or path.stat().st_size,'managed')
        meta=next(x for x in local_metadata() if x['name']==table)
        return {'ok':True,'table':table,'rows':rows,'columns':[str(c) for c in df.columns],'metadata':meta,'storage':meta.get('storage'),'sourceType':'file','sheet':selected,'sheets':sheet_names,'displayName':friendly_table_name(selected)}
    if ext=='.xml':
        df=_xml_frame(path)
        rows=write_dataframe(df,table,source_bytes or path.stat().st_size,'managed')
        meta=next(x for x in local_metadata() if x['name']==table)
        return {'ok':True,'table':table,'rows':rows,'columns':[str(c) for c in df.columns],'metadata':meta,'storage':meta.get('storage'),'sourceType':'file','displayName':friendly_table_name(table)}
    raise ValueError('Supported file types are CSV, TSV/TXT, XLSX/XLS, JSON/JSONL, Parquet and XML.')

def import_file(filename:str,raw:bytes,sheet:str|None=None):
    suffix=Path(filename).suffix or '.csv'
    with tempfile.NamedTemporaryFile(delete=False,suffix=suffix) as f:
        f.write(raw);tmp=Path(f.name)
    try:return import_file_path(filename,tmp,sheet,len(raw))
    finally:
        try:tmp.unlink()
        except Exception:pass

def _google_sheet_csv(url:str):
    """Convert a normal Google Sheets URL to a public CSV export URL.

    Keep the worksheet gid only when the shared link actually contains one.
    Older code forced gid=0, but Google worksheet ids are not guaranteed to be 0.
    Links copied from Google can also place gid in the URL fragment (#gid=...).
    """
    if '/spreadsheets/d/' in url:
        m=re.search(r'/spreadsheets/d/([^/]+)',url)
        sid=m.group(1) if m else None
        parsed=urllib.parse.urlparse(url)
        query=urllib.parse.parse_qs(parsed.query)
        fragment=urllib.parse.parse_qs(parsed.fragment)
        gid=(query.get('gid') or fragment.get('gid') or [None])[0]
        if sid:
            export=f'https://docs.google.com/spreadsheets/d/{sid}/export?format=csv'
            if gid is not None and str(gid).strip():
                export+=f'&gid={urllib.parse.quote(str(gid).strip(),safe="")}'
            return export
    return url

def _validate_remote_url(url:str):
    """Block local/private-network SSRF by default for cloud/API imports."""
    parsed=urllib.parse.urlparse(url)
    if parsed.scheme not in ('https','http') or not parsed.hostname:
        raise ValueError('Only HTTP/HTTPS URLs are allowed.')
    if parsed.scheme!='https' and __import__('os').environ.get('VTAB_ALLOW_INSECURE_HTTP','0')!='1':
        raise ValueError('HTTPS is required for remote data sources. Set VTAB_ALLOW_INSECURE_HTTP=1 only for trusted development endpoints.')
    if __import__('os').environ.get('VTAB_ALLOW_PRIVATE_REMOTE','0')=='1':
        return
    try:
        infos=socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme=='https' else 80), type=socket.SOCK_STREAM)
        for info in infos:
            ip=ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast or ip.is_reserved:
                raise ValueError('Private, loopback, link-local and reserved network destinations are blocked for remote imports.')
    except socket.gaierror as exc:
        raise ValueError(f'Could not resolve remote host: {exc}')

def _remote_request(url:str,access_token:str|None=None):
    headers={'User-Agent':'VTAB-Reporting-Studio/4.2','Accept':'*/*'}
    if access_token:headers['Authorization']='Bearer '+access_token.strip()
    return urllib.request.Request(url,headers=headers)

def _microsoft_download_url(source_type:str,url:str,access_token:str|None):
    if 'graph.microsoft.com/' in url:return url
    if access_token and ('sharepoint.com/' in url or '1drv.ms/' in url or 'onedrive.live.com/' in url):
        encoded=base64.urlsafe_b64encode(url.encode()).decode().rstrip('=')
        return f'https://graph.microsoft.com/v1.0/shares/u!{encoded}/driveItem/content'
    return url

def list_google_sheets(url:str,access_token:str|None=None):
    """Return worksheet tab names from a Google Spreadsheet.
    With a token: calls the Sheets API for full metadata.
    Without a token: downloads the XLSX export and reads sheet names from the workbook."""
    match=re.search(r'/spreadsheets/d/([^/]+)',url)
    if not match:raise ValueError('Could not read the Google spreadsheet ID from the URL.')
    sid=match.group(1).split('/')[0].split('?')[0]
    if access_token:
        api_url=f'https://sheets.googleapis.com/v4/spreadsheets/{sid}?fields=sheets.properties'
        _validate_remote_url(api_url)
        with urllib.request.urlopen(_remote_request(api_url,access_token),timeout=30) as resp:
            payload=json.load(resp)
        sheets=payload.get('sheets') or []
        return [{'title':s['properties']['title'],'sheetId':s['properties']['sheetId'],'index':s['properties']['index']} for s in sheets]
    # No token: download the public XLSX export and read worksheet names via openpyxl
    export_url=f'https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx'
    _validate_remote_url(export_url)
    with urllib.request.urlopen(_remote_request(export_url),timeout=60) as resp:
        xlsx_data=resp.read()
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(xlsx_data),read_only=True,data_only=True)
    names=list(wb.sheetnames);wb.close()
    return [{'title':n,'sheetId':i,'index':i} for i,n in enumerate(names)]

def _download_public_xlsx(url:str)->bytes:
    """Download a public Google Sheet as XLSX bytes, respecting the size limit."""
    match=re.search(r'/spreadsheets/d/([^/]+)',url)
    if not match:raise ValueError('Could not read the Google spreadsheet ID from the URL.')
    sid=match.group(1).split('/')[0].split('?')[0]
    export_url=f'https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx'
    _validate_remote_url(export_url)
    max_bytes=int(os.environ.get('VTAB_MAX_REMOTE_MB','2048'))*1024*1024
    chunks=[];total=0
    with urllib.request.urlopen(_remote_request(export_url),timeout=60) as resp:
        while True:
            chunk=resp.read(8*1024*1024)
            if not chunk:break
            chunks.append(chunk);total+=len(chunk)
            if total>max_bytes:raise ValueError(f'Google Sheets export exceeds VTAB_MAX_REMOTE_MB ({max_bytes//1024//1024} MB).')
    return b''.join(chunks)

def _build_sheet_range(sheet_title:str,col_range:str|None)->str:
    """Prefix column range with sheet title. If col_range already contains '!', return as-is."""
    col_range=(col_range or 'A:ZZ').strip()
    if '!' in col_range:return col_range
    # Escape sheet title for Sheets API (wrap in single quotes if needed)
    safe_title=sheet_title.replace("'","''")
    return f"'{safe_title}'!{col_range}"

def import_cloud(source_type:str,url:str,name:str|None=None,access_token:str|None=None,sheet_range:str|None=None):
    source_type=(source_type or '').lower()
    if source_type=='google_sheets' and access_token and '/spreadsheets/d/' in url:
        match=re.search(r'/spreadsheets/d/([^/]+)',url)
        if not match:raise ValueError('Could not read the Google spreadsheet ID from the URL.')
        # Build a well-formed range: if sheet_range already has '!', use as-is;
        # otherwise default to first sheet prefix to avoid the 400 Bad Request.
        raw_range=sheet_range or 'A:ZZ'
        if '!' not in raw_range:
            # Fetch first sheet name to produce a valid prefixed range
            try:
                first=list_google_sheets(url,access_token)
                if first and not first[0].get('fallback'):
                    raw_range=_build_sheet_range(first[0]['title'],raw_range)
            except Exception:
                pass  # If metadata fetch fails, try without prefix (may still 400)
        cell_range=urllib.parse.quote(raw_range,safe='!:$')
        api_url=f'https://sheets.googleapis.com/v4/spreadsheets/{match.group(1)}/values/{cell_range}?majorDimension=ROWS'
        _validate_remote_url(api_url)
        with urllib.request.urlopen(_remote_request(api_url,access_token),timeout=60) as resp:
            payload=json.load(resp)
        values=payload.get('values') or []
        if not values:raise ValueError('Google Sheets returned no cells for the selected range.')
        headers=[str(x) or f'Column_{i+1}' for i,x in enumerate(values[0])]
        rows=[list(row)+[None]*(len(headers)-len(row)) for row in values[1:]]
        df=pd.DataFrame([row[:len(headers)] for row in rows],columns=headers)
        table=safe_table_name(name or 'GoogleSheet')
        count=write_dataframe(df,table,None,'managed')
        meta=next(x for x in local_metadata() if x['name']==table)
        return {'ok':True,'table':table,'rows':count,'columns':headers,'metadata':meta,'storage':meta.get('storage'),'sourceType':'cloud','displayName':friendly_table_name(table)}
    # Public Google Sheet with a specific named sheet: download XLSX and read that sheet
    if source_type=='google_sheets' and not access_token and '/spreadsheets/d/' in url and sheet_range and '!' in sheet_range:
        sheet_name=sheet_range.split('!')[0].strip("' ")
        xlsx_data=_download_public_xlsx(url)
        df=pd.read_excel(io.BytesIO(xlsx_data),sheet_name=sheet_name,engine='openpyxl')
        table=safe_table_name(name or sheet_name or 'GoogleSheet')
        count=write_dataframe(df,table,None,'managed')
        meta=next(x for x in local_metadata() if x['name']==table)
        return {'ok':True,'table':table,'rows':count,'columns':list(df.columns),'metadata':meta,'storage':meta.get('storage'),'sourceType':'cloud','displayName':friendly_table_name(table)}
    if source_type=='google_sheets':url=_google_sheet_csv(url)
    if source_type in ('sharepoint','onedrive'):url=_microsoft_download_url(source_type,url,access_token)
    _validate_remote_url(url)
    req=_remote_request(url,access_token)
    with urllib.request.urlopen(req,timeout=60) as resp:
        ctype=(resp.headers.get('Content-Type') or '').lower()
        disposition=resp.headers.get('Content-Disposition') or ''
        file_match=re.search(r'filename\*?=(?:UTF-8\'\')?["\']?([^"\';]+)',disposition,re.I)
        remote_name=urllib.parse.unquote(file_match.group(1)) if file_match else ''
        suffix=Path(remote_name).suffix.lower()
        if not suffix:suffix='.json' if 'json' in ctype else '.xlsx' if 'spreadsheet' in ctype or 'excel' in ctype else '.parquet' if 'parquet' in ctype else '.xml' if 'xml' in ctype else '.csv'
        filename=name or ('GoogleSheet'+suffix if source_type=='google_sheets' else 'CloudData'+suffix)
        if not Path(filename).suffix:filename+=suffix
        max_bytes=int(os.environ.get('VTAB_MAX_REMOTE_MB','2048'))*1024*1024
        with tempfile.NamedTemporaryFile(delete=False,suffix=suffix) as f:
            total=0
            while True:
                chunk=resp.read(8*1024*1024)
                if not chunk:break
                f.write(chunk);total+=len(chunk)
                if total>max_bytes:raise ValueError(f'Remote source exceeds VTAB_MAX_REMOTE_MB ({max_bytes//1024//1024} MB).')
            tmp=Path(f.name)
    try:
        result=import_file_path(filename,tmp,None,total)
        # Override sourceType to 'cloud' for remote API/cloud imports
        result['sourceType']='cloud'
        return result
    finally:
        try:tmp.unlink()
        except Exception:pass

def _db_config(cfg:dict,typ:str):
    host=cfg.get('host') or cfg.get('server')
    database=cfg.get('database')
    if not host or not database:raise ValueError('Server/host and database are required.')
    return host,database

def _driver_error(label:str,requirement:str,exc:Exception):
    raise ValueError(f'{label} connector driver is unavailable. Install/package {requirement}. Details: {exc}') from exc

def _is_loopback_host(host:str)->bool:
    value=str(host or '').strip().lower().strip('[]')
    if value=='localhost':
        return True
    try:
        return ipaddress.ip_address(value).is_loopback
    except ValueError:
        return False


def _postgres_connect(cfg:dict):
    try:import psycopg
    except Exception as exc:_driver_error('PostgreSQL/Redshift','psycopg[binary]',exc)
    connection_string=str(cfg.get('connectionString') or '').strip()
    if connection_string:return psycopg.connect(connection_string,connect_timeout=15)
    host,database=_db_config(cfg,'postgresql')
    sslmode=str(cfg.get('sslmode') or 'prefer').strip().lower()
    # Local PostgreSQL installations commonly run without TLS.  Keep cloud/private
    # hosts strict when the user requires SSL, but allow libpq/psycopg to try SSL
    # and safely fall back to a non-SSL loopback connection for localhost only.
    if sslmode=='require' and _is_loopback_host(host):
        sslmode='prefer'
    return psycopg.connect(host=host,port=int(cfg.get('port') or 5432),dbname=database,user=cfg.get('user') or cfg.get('username'),password=cfg.get('password'),sslmode=sslmode,connect_timeout=15)

def _sqlserver_connection_string(cfg:dict)->str:
    connection_string=str(cfg.get('connectionString') or '').strip()
    if connection_string:return connection_string
    host,database=_db_config(cfg,'sqlserver')
    driver=cfg.get('driver') or 'ODBC Driver 18 for SQL Server'
    server=host+(f',{int(cfg["port"])}' if cfg.get('port') else '')
    auth='Trusted_Connection=yes' if cfg.get('trustedConnection') else f'UID={cfg.get("user") or cfg.get("username") or ""};PWD={cfg.get("password","")}'
    # The Services UI stores the encryption checkbox as sslmode=require/disable.
    # Honour an explicit `encrypt` value first, otherwise translate sslmode while
    # keeping the secure default (Encrypt=yes) for existing configurations.
    explicit_encrypt=cfg.get('encrypt')
    if explicit_encrypt is None:
        encrypt_enabled=str(cfg.get('sslmode') or 'require').strip().lower() not in ('disable','false','0','no')
    else:
        encrypt_enabled=str(explicit_encrypt).strip().lower() not in ('disable','false','0','no')
    encrypt='yes' if encrypt_enabled else 'no'
    trust='yes' if cfg.get('trustServerCertificate') else 'no'
    return f'DRIVER={{{driver}}};SERVER={server};DATABASE={database};{auth};Encrypt={encrypt};TrustServerCertificate={trust}'

def _odbc_connect(cfg:dict):
    try:import pyodbc
    except Exception as exc:_driver_error('ODBC','pyodbc plus the database vendor ODBC driver',exc)
    connection_string=str(cfg.get('connectionString') or '').strip()
    if not connection_string:raise ValueError('An ODBC connection string is required.')
    return pyodbc.connect(connection_string,timeout=15)

def _sqlserver_connect(cfg:dict):
    try:import pyodbc
    except Exception as exc:_driver_error('SQL Server','pyodbc and Microsoft ODBC Driver 18 for SQL Server',exc)
    return pyodbc.connect(_sqlserver_connection_string(cfg),timeout=15)

def _mysql_connect(cfg:dict):
    try:import pymysql
    except Exception as exc:_driver_error('MySQL/MariaDB','pymysql',exc)
    values=dict(cfg)
    connection_string=str(values.get('connectionString') or '').strip()
    if connection_string:
        parsed=urllib.parse.urlparse(connection_string)
        if parsed.scheme not in ('mysql','mariadb'):raise ValueError('MySQL connection string must start with mysql:// or mariadb://.')
        values.update({'host':parsed.hostname,'port':parsed.port,'user':urllib.parse.unquote(parsed.username or ''),'password':urllib.parse.unquote(parsed.password or ''),'database':parsed.path.lstrip('/')})
    host,database=_db_config(values,'mysql')
    ssl={'ssl':{}} if str(values.get('sslmode') or '').lower() not in ('','disable','false','0') else {}
    return pymysql.connect(host=host,port=int(values.get('port') or 3306),database=database,user=values.get('user') or values.get('username'),password=values.get('password') or '',connect_timeout=15,read_timeout=120,write_timeout=120,charset='utf8mb4',**ssl)

def _sqlite_connect(cfg:dict):
    import sqlite3
    path=Path(str(cfg.get('path') or '').strip()).expanduser()
    if not path.is_file():raise ValueError(f'SQLite database file was not found: {path}')
    return sqlite3.connect(f'file:{path.resolve().as_posix()}?mode=ro',uri=True)

def _database_family(source_type:str)->str:
    aliases={'postgres':'postgresql','redshift':'postgresql','mariadb':'mysql','synapse':'sqlserver','fabric_warehouse':'sqlserver','access':'odbc'}
    return aliases.get((source_type or '').lower(),(source_type or '').lower())

def _database_connect(source_type:str,cfg:dict):
    if (source_type or '').lower()=='access' and not cfg.get('connectionString'):
        path=str(cfg.get('path') or '').strip()
        if not path:raise ValueError('An Access database path or ODBC connection string is required.')
        cfg={**cfg,'connectionString':f'DRIVER={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={path}'}
    family=_database_family(source_type)
    if family=='postgresql':return _postgres_connect(cfg)
    if family=='sqlserver':return _sqlserver_connect(cfg)
    if family=='mysql':return _mysql_connect(cfg)
    if family=='sqlite':return _sqlite_connect(cfg)
    if family=='odbc':return _odbc_connect(cfg)
    raise ValueError(f'{source_type} does not currently provide a database navigator. Use a supported relational connector or install its VTAB plugin.')

def test_connection(payload):
    typ=(payload.get('type') or '').lower()
    cfg=payload.get('config') or {}
    if typ=='demo':return {'ok':True,'message':'Local DuckDB columnar workspace is ready.'}
    if typ in ('postgresql','postgres','redshift','sqlserver','synapse','fabric_warehouse','mysql','mariadb','sqlite','odbc','access'):
        conn=_database_connect(typ,cfg)
        try:
            cur=conn.cursor();cur.execute('SELECT 1');cur.fetchone();cur.close()
        finally:conn.close()
        return {'ok':True,'message':f'{next((x["name"] for x in SOURCE_CATALOG if x["id"]==typ),typ)} connection successful.'}
    if typ in ('oracle','db2','snowflake','databricks','bigquery'):
        raise ValueError(f'{typ} requires a dedicated VTAB connector plugin; it is not a working connector in this build.')
    if typ in ('rest','odata','graphql'):
        url=cfg.get('url','') or cfg.get('endpoint','') or cfg.get('base_url','')
        if not url:raise ValueError('An HTTPS endpoint is required.')
        _validate_remote_url(url)
        headers={str(k):str(v) for k,v in (cfg.get('headers') or {}).items()}
        token=cfg.get('api_key') or cfg.get('token') or cfg.get('access_token')
        if token and 'Authorization' not in headers:
            headers['Authorization']=f'Bearer {token}'
        request=urllib.request.Request(url,headers=headers,method='GET')
        try:
            with urllib.request.urlopen(request,timeout=max(1,min(int(cfg.get('timeout') or 30),120))) as response:
                response.read(1)
                return {'ok':True,'message':f'{typ.upper()} connection and authentication successful (HTTP {response.status}).'}
        except urllib.error.HTTPError as error:
            raise ValueError(f'Endpoint authentication/open failed with HTTP {error.code}: {error.reason}') from error
        except Exception as error:
            raise ValueError(f'Unable to open endpoint: {error}') from error
    if typ in ('sharepoint','onedrive','google_sheets','azure_blob','adls_gen2','s3','gcs','salesforce','dynamics365','servicenow','jira','github'):
        url=cfg.get('url','') or cfg.get('endpoint','')
        if not url:raise ValueError('An HTTPS endpoint/shared/export URL is required.')
        _validate_remote_url(url)
        return {'ok':True,'message':'Secure endpoint configuration accepted. Authentication is validated when the connector executes.'}
    raise ValueError('Unsupported connector')

def _safe_select(query:str):
    value=(query or '').strip().rstrip(';')
    if not re.match(r'^(select|with)\b',value,re.I):raise ValueError('Only SELECT/CTE queries are allowed for database imports.')
    if re.search(r'\b(insert|update|delete|drop|alter|truncate|merge|execute|exec|grant|revoke)\b',value,re.I):raise ValueError('The database import query must be read-only.')
    return value

def list_database_tables(source_type:str,config:dict):
    source_type=(source_type or '').lower()
    family=_database_family(source_type)
    if family=='postgresql':
        with _postgres_connect(config) as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT table_schema,table_name FROM information_schema.tables WHERE table_schema NOT IN ('information_schema','pg_catalog') AND table_type='BASE TABLE' ORDER BY table_schema,table_name")
                return [{'schema':s,'name':n,'qualifiedName':f'{s}.{n}','query':f'SELECT * FROM {qident(s)}.{qident(n)}'} for s,n in cur.fetchall()]
    if family=='sqlserver':
        conn=_sqlserver_connect(config)
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT table_schema,table_name FROM information_schema.tables WHERE table_type='BASE TABLE' ORDER BY table_schema,table_name")
                return [{'schema':s,'name':n,'qualifiedName':f'{s}.{n}','query':f'SELECT * FROM {bident(s)}.{bident(n)}'} for s,n in cur.fetchall()]
        finally:conn.close()
    if family=='mysql':
        conn=_mysql_connect(config)
        try:
            with conn.cursor() as cur:
                cur.execute('SELECT DATABASE()');database=cur.fetchone()[0]
                if not database:raise ValueError('MySQL did not select a database. Add the database name to the connection settings.')
                cur.execute("SELECT table_schema,table_name FROM information_schema.tables WHERE table_schema=%s AND table_type='BASE TABLE' ORDER BY table_name",(database,))
                return [{'schema':s,'name':n,'qualifiedName':f'{s}.{n}','query':f'SELECT * FROM {tickident(s)}.{tickident(n)}'} for s,n in cur.fetchall()]
        finally:conn.close()
    if family=='sqlite':
        conn=_sqlite_connect(config)
        try:
            names=[r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name").fetchall()]
            return [{'schema':'main','name':n,'qualifiedName':n,'query':f'SELECT * FROM {qident(n)}'} for n in names]
        finally:conn.close()
    if family=='odbc':
        conn=_odbc_connect(config)
        try:
            rows=[]
            for row in conn.cursor().tables(tableType='TABLE'):
                schema=str(row.table_schem or '').strip();name=str(row.table_name);qualified=f'{schema}.{name}' if schema else name
                query='SELECT * FROM '+'.'.join(bident(x) for x in ([schema,name] if schema else [name]))
                rows.append({'schema':schema,'name':name,'qualifiedName':qualified,'query':query})
            return sorted(rows,key=lambda x:x['qualifiedName'].casefold())
        finally:conn.close()
    raise ValueError(f'Database table discovery is not implemented for {source_type}.')

def qident(value:str)->str:return '"'+str(value).replace('"','""')+'"'
def bident(value:str)->str:return '['+str(value).replace(']',']]')+']'
def tickident(value:str)->str:return '`'+str(value).replace('`','``')+'`'

def import_database(source_type:str,config:dict,query:str,name:str|None=None):
    source_type=(source_type or '').lower();family=_database_family(source_type);query=_safe_select(query)
    if family=='postgresql':
        with _postgres_connect(config) as conn:
            df=pd.read_sql_query(query,conn)
    elif family in ('sqlserver','odbc'):
        conn=_sqlserver_connect(config) if family=='sqlserver' else _odbc_connect(config)
        try:df=pd.read_sql_query(query,conn)
        finally:conn.close()
    elif family=='mysql':
        conn=_mysql_connect(config)
        try:df=pd.read_sql_query(query,conn)
        finally:conn.close()
    elif family=='sqlite':
        conn=_sqlite_connect(config)
        try:df=pd.read_sql_query(query,conn)
        finally:conn.close()
    else:raise ValueError(f'Database import is not implemented for {source_type}.')
    table=safe_table_name(name or f'{source_type}_query')
    rows=write_dataframe(df,table,None,'managed');meta=next(x for x in local_metadata() if x['name']==table)
    return {'ok':True,'table':table,'rows':rows,'columns':[str(c) for c in df.columns],'metadata':meta,'storage':meta.get('storage'),'sourceType':'database','displayName':friendly_table_name(table)}

SOURCE_CATALOG=[
 # Databases
 {'id':'sqlserver','name':'SQL Server / Azure SQL','category':'Database','mode':'Import','status':'native'},
 {'id':'postgresql','name':'PostgreSQL','category':'Database','mode':'Import','status':'native'},
 {'id':'mysql','name':'MySQL','category':'Database','mode':'Import','status':'native'},
 {'id':'mariadb','name':'MariaDB','category':'Database','mode':'Import','status':'native'},
 {'id':'oracle','name':'Oracle Database','category':'Database','mode':'DirectQuery / Import','status':'plugin'},
 {'id':'db2','name':'IBM Db2','category':'Database','mode':'DirectQuery / Import','status':'plugin'},
 {'id':'sqlite','name':'SQLite','category':'Database','mode':'Import','status':'native'},
 {'id':'access','name':'Microsoft Access','category':'Database','mode':'ODBC Import','status':'driver'},
 {'id':'odbc','name':'ODBC','category':'Database','mode':'Import','status':'driver'},
 {'id':'jdbc','name':'JDBC','category':'Database','mode':'Import','status':'plugin'},
 # Warehouses / lakehouses
 {'id':'snowflake','name':'Snowflake','category':'Cloud Warehouse','mode':'DirectQuery / Import','status':'plugin'},
 {'id':'databricks','name':'Databricks SQL / Lakehouse','category':'Cloud Warehouse','mode':'DirectQuery / Import','status':'plugin'},
 {'id':'bigquery','name':'Google BigQuery','category':'Cloud Warehouse','mode':'DirectQuery / Import','status':'plugin'},
 {'id':'redshift','name':'Amazon Redshift','category':'Cloud Warehouse','mode':'Import','status':'native'},
 {'id':'synapse','name':'Azure Synapse Analytics','category':'Cloud Warehouse','mode':'Import','status':'native'},
 {'id':'fabric_warehouse','name':'Microsoft Fabric Warehouse / Lakehouse','category':'Cloud Warehouse','mode':'Import','status':'native'},
 # Files
 {'id':'csv','name':'CSV','category':'File','mode':'Columnar Import','status':'native'},
 {'id':'tsv','name':'Text / TSV','category':'File','mode':'Columnar Import','status':'native'},
 {'id':'excel','name':'Excel (XLSX/XLS)','category':'File','mode':'Columnar Import','status':'native'},
 {'id':'json','name':'JSON / JSONL','category':'File','mode':'Columnar Import','status':'native'},
 {'id':'parquet','name':'Parquet','category':'File','mode':'Zero-copy / Import','status':'native'},
 {'id':'xml','name':'XML','category':'File','mode':'Import','status':'native'},
 {'id':'folder','name':'Folder / Append','category':'File','mode':'Batch Import','status':'native'},
 # Cloud storage/files
 {'id':'sharepoint','name':'SharePoint','category':'Cloud Storage','mode':'Import','status':'native'},
 {'id':'onedrive','name':'OneDrive','category':'Cloud Storage','mode':'Import','status':'native'},
 {'id':'google_sheets','name':'Google Sheets','category':'Cloud Storage','mode':'Import','status':'native'},
 {'id':'azure_blob','name':'Azure Blob Storage','category':'Cloud Storage','mode':'Import','status':'endpoint'},
 {'id':'adls_gen2','name':'Azure Data Lake Storage Gen2','category':'Cloud Storage','mode':'Import','status':'endpoint'},
 {'id':'s3','name':'Amazon S3','category':'Cloud Storage','mode':'Import','status':'endpoint'},
 {'id':'gcs','name':'Google Cloud Storage','category':'Cloud Storage','mode':'Import','status':'endpoint'},
 # APIs / SaaS
 {'id':'rest','name':'REST API','category':'API / SaaS','mode':'Import / Refresh','status':'native'},
 {'id':'odata','name':'OData Feed','category':'API / SaaS','mode':'Import / Refresh','status':'native'},
 {'id':'graphql','name':'GraphQL API','category':'API / SaaS','mode':'Import / Refresh','status':'endpoint'},
 {'id':'salesforce','name':'Salesforce','category':'API / SaaS','mode':'Import / Refresh','status':'plugin'},
 {'id':'dynamics365','name':'Dynamics 365 / Dataverse','category':'API / SaaS','mode':'Import / Refresh','status':'plugin'},
 {'id':'servicenow','name':'ServiceNow','category':'API / SaaS','mode':'Import / Refresh','status':'plugin'},
 {'id':'jira','name':'Jira','category':'API / SaaS','mode':'Import / Refresh','status':'plugin'},
 {'id':'github','name':'GitHub','category':'API / SaaS','mode':'Import / Refresh','status':'plugin'},
]
